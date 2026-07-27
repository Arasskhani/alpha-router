"""Render chat assistant content (markdown) to a Word .docx file.

Security model:
- The caller is an authenticated user exporting their own chat content.
- ``python-docx`` only constructs an OpenXML package; it does not execute
  embedded code, fetch network resources, or evaluate macros. There is no
  SSRF/XSS surface (unlike an HTML/PDF render path).
- No content length cap is enforced (per product decision).
- Markdown is parsed structurally (not via HTML), so raw HTML in the content
  is treated as plain text and never rendered as markup.
"""

from __future__ import annotations

import io
import re
from typing import Iterable

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


class ChatExportError(Exception):
    """Chat DOCX render failed."""


# ---------------------------------------------------------------------------
# RTL helpers
# ---------------------------------------------------------------------------


def _set_paragraph_rtl(paragraph: Paragraph) -> None:
    """Mark a paragraph as right-to-left (Arabic/Persian)."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)


def _set_run_rtl(run) -> None:
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    rPr.append(rtl)


# ---------------------------------------------------------------------------
# Inline formatting: **bold**, *italic* / _italic_, `code`, [text](url),
# ~~strike~~. Everything else is literal text.
# ---------------------------------------------------------------------------

_INLINE_RE = re.compile(
    r"""
    (?P<code>`[^`\n]+`)
    | (?P<bold>\*\*[^*\n]+\*\*)
    | (?P<italic>\*[^*\n]+\*|_[^_\n]+_)
    | (?P<strike>~~[^~\n]+~~)
    | (?P<link>\[[^\]\n]*\]\([^)\n]+\))
    """,
    re.VERBOSE,
)


def _add_inline_runs(paragraph: Paragraph, text: str, *, rtl: bool) -> None:
    """Add runs to ``paragraph`` for ``text`` with inline markdown formatting."""
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            if rtl:
                _set_run_rtl(run)
        token = m.group(0)
        if m.group("code"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10.5)
        elif m.group("bold"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif m.group("italic"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif m.group("strike"):
            run = paragraph.add_run(token[2:-2])
            run.font.strike = True
        elif m.group("link"):
            inner = re.match(r"\[([^\]]*)\]\(([^)]+)\)", token)
            label = inner.group(1) if inner else token
            run = paragraph.add_run(label)
            run.font.color.rgb = RGBColor(0x09, 0x69, 0xDA)
            run.underline = True
        if rtl and "run" in locals():
            _set_run_rtl(run)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        if rtl:
            _set_run_rtl(run)


# ---------------------------------------------------------------------------
# Block-level markdown parser -> docx
# ---------------------------------------------------------------------------


_FENCE_RE = re.compile(r"^```(.*)$")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_HR_RE = re.compile(r"^\s*([-*_])\1{2,}\s*$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_ULIST_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
_OLIST_RE = re.compile(r"^\s*\d+\.\s+(.*)$")
_QUOTE_RE = re.compile(r"^\s*>\s?(.*)$")


def _is_table_row(line: str) -> bool:
    return "|" in line and line.strip().startswith("|") or line.strip().endswith("|")


def _split_pipe_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _add_heading(doc: DocxDocument, level: int, text: str) -> None:
    para = doc.add_heading(level=min(level, 6))
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_rtl(para)
    _add_inline_runs(para, text, rtl=True)


def _add_paragraph(doc: DocxDocument, text: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_rtl(para)
    _add_inline_runs(para, text, rtl=True)


def _add_code_block(doc: DocxDocument, code: str, lang: str = "") -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    # light shading via paragraph borders is complex; keep simple.


def _add_list_item(doc: DocxDocument, text: str, *, ordered: bool) -> None:
    style = "List Number" if ordered else "List Bullet"
    try:
        para = doc.add_paragraph(style=style)
    except KeyError:
        para = doc.add_paragraph(("1. " if ordered else "- ") + text)
        _add_inline_runs(para, "", rtl=True)
        return
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_rtl(para)
    _add_inline_runs(para, text, rtl=True)


def _add_table(doc: DocxDocument, header: list[str], rows: list[list[str]]) -> None:
    cols = max(len(header), *(len(r) for r in rows)) if rows else len(header)
    table: Table = doc.add_table(rows=1 + len(rows), cols=cols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    # header
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_paragraph_rtl(para)
        run = para.add_run(cell_text)
        run.bold = True
        _set_run_rtl(run)
    # body
    for r, row in enumerate(rows, start=1):
        for c, cell_text in enumerate(row):
            if c >= cols:
                break
            cell = table.rows[r].cells[c]
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_paragraph_rtl(para)
            _add_inline_runs(para, cell_text, rtl=True)


def _add_quote(doc: DocxDocument, text: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_rtl(para)
    para.paragraph_format.left_indent = Pt(18)
    run = para.add_run("❖ ")
    _set_run_rtl(run)
    _add_inline_runs(para, text, rtl=True)


def _add_hr(doc: DocxDocument) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run("— — — — — — — — — —")
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


def _render_markdown_to_docx(doc: DocxDocument, content: str) -> None:
    lines = content.split("\n")
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # fenced code block
        fence = _FENCE_RE.match(line)
        if fence:
            lang = fence.group(1).strip()
            buf: list[str] = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            _add_code_block(doc, "\n".join(buf), lang)
            continue

        # heading
        h = _HEADING_RE.match(line)
        if h:
            level = len(h.group(1))
            _add_heading(doc, level, h.group(2).strip())
            i += 1
            continue

        # horizontal rule
        if _HR_RE.match(line):
            _add_hr(doc)
            i += 1
            continue

        # table: header + separator
        if _is_table_row(line) and i + 1 < n and _TABLE_SEP_RE.match(lines[i + 1]):
            header = _split_pipe_row(line)
            i += 2  # skip header + separator
            rows: list[list[str]] = []
            while i < n and _is_table_row(lines[i]) and not _TABLE_SEP_RE.match(lines[i]):
                rows.append(_split_pipe_row(lines[i]))
                i += 1
            _add_table(doc, header, rows)
            continue

        # blockquote (consecutive)
        if _QUOTE_RE.match(line):
            buf2: list[str] = []
            while i < n and _QUOTE_RE.match(lines[i]):
                buf2.append(_QUOTE_RE.match(lines[i]).group(1))
                i += 1
            _add_quote(doc, " ".join(buf2))
            continue

        # unordered list
        if _ULIST_RE.match(line):
            while i < n and _ULIST_RE.match(lines[i]):
                _add_list_item(doc, _ULIST_RE.match(lines[i]).group(1), ordered=False)
                i += 1
            continue

        # ordered list
        if _OLIST_RE.match(line):
            while i < n and _OLIST_RE.match(lines[i]):
                _add_list_item(doc, _OLIST_RE.match(lines[i]).group(1), ordered=True)
                i += 1
            continue

        # blank line
        if not line.strip():
            i += 1
            continue

        # paragraph (gather consecutive non-empty, non-special lines)
        buf3: list[str] = [line]
        i += 1
        while i < n and lines[i].strip() and not _FENCE_RE.match(lines[i]) and not _HEADING_RE.match(lines[i]) \
                and not _HR_RE.match(lines[i]) and not _ULIST_RE.match(lines[i]) \
                and not _OLIST_RE.match(lines[i]) and not _QUOTE_RE.match(lines[i]) \
                and not (_is_table_row(lines[i]) and i + 1 < n and _TABLE_SEP_RE.match(lines[i + 1])):
            buf3.append(lines[i])
            i += 1
        _add_paragraph(doc, " ".join(s.strip() for s in buf3))


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def render_chat_docx(*, content: str, title: str | None = None) -> bytes:
    """Render markdown ``content`` to a .docx byte string."""
    if content is None or not content.strip():
        raise ChatExportError("content must not be empty")
    doc = Document()
    # base font friendly to Persian
    style = doc.styles["Normal"]
    style.font.name = "Tahoma"
    style.font.size = Pt(11)

    if title and title.strip():
        h = doc.add_heading(title.strip(), level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_paragraph_rtl(h)

    _render_markdown_to_docx(doc, content)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
