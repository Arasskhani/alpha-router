"""Strict validation and bounded text extraction for Knowledge documents."""

from __future__ import annotations

import csv
import json
import re
import unicodedata
from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.config import get_settings
from app.core.archive_safety import ArchiveSafetyError, validate_ooxml_archive


class UnsafeDocumentError(ValueError):
    pass


@dataclass(frozen=True)
class ValidatedDocument:
    file_name: str
    extension: str
    mime_type: str
    size_bytes: int
    format_name: str


@dataclass(frozen=True)
class ParsedSegment:
    text: str
    page_number: int | None = None
    section: str | None = None


@dataclass(frozen=True)
class ParsedDocument:
    segments: tuple[ParsedSegment, ...]
    character_count: int
    metadata: dict
    security_flags: tuple[str, ...]


_FORMAT_BY_EXTENSION = {
    ".pdf": ("pdf", "application/pdf"),
    ".docx": (
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
    ".pptx": (
        "pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ),
    ".xlsx": (
        "xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ),
    ".txt": ("text", "text/plain"),
    ".md": ("markdown", "text/markdown"),
    ".markdown": ("markdown", "text/markdown"),
    ".html": ("html", "text/html"),
    ".htm": ("html", "text/html"),
    ".csv": ("csv", "text/csv"),
    ".json": ("json", "application/json"),
}

_GENERIC_MIMES = {
    "",
    "application/octet-stream",
    "binary/octet-stream",
    "application/zip",
}

_BIDI_CONTROLS = {
    "\u202a",
    "\u202b",
    "\u202c",
    "\u202d",
    "\u202e",
    "\u2066",
    "\u2067",
    "\u2068",
    "\u2069",
}

_PROMPT_INJECTION_PATTERNS = (
    re.compile(r"\bignore\s+(all\s+)?previous\s+instructions\b", re.I),
    re.compile(r"\b(system|developer)\s+prompt\b", re.I),
    re.compile(r"<\s*/?\s*(system|assistant|developer)\s*>", re.I),
    re.compile(r"\bdo\s+not\s+follow\s+the\s+user\b", re.I),
)


def sanitize_document_filename(file_name: str) -> str:
    base = Path((file_name or "").replace("\\", "/")).name
    base = "".join(
        character
        for character in unicodedata.normalize("NFC", base)
        if character not in _BIDI_CONTROLS and (character.isprintable() or character in {" ", "\t"})
    )
    base = re.sub(r"\s+", " ", base).strip(" .")
    if not base or len(base) > 240:
        raise UnsafeDocumentError("Document filename is invalid")
    return base


def _validate_ooxml_archive(data: bytes, extension: str) -> None:
    try:
        validate_ooxml_archive(data, extension)
    except ArchiveSafetyError as exc:
        raise UnsafeDocumentError(str(exc)) from exc


def _validate_pdf(data: bytes) -> None:
    settings = get_settings()
    try:
        reader = PdfReader(BytesIO(data), strict=True)
    except Exception as exc:
        raise UnsafeDocumentError("PDF structure is invalid") from exc
    if reader.is_encrypted:
        raise UnsafeDocumentError("Encrypted PDF documents are not supported")
    if len(reader.pages) > settings.knowledge_max_pdf_pages:
        raise UnsafeDocumentError("PDF contains too many pages")
    root = reader.trailer.get("/Root") or {}
    if hasattr(root, "get_object"):
        root = root.get_object()
    names = root.get("/Names") or {}
    if hasattr(names, "get_object"):
        names = names.get_object()
    if any(key in names for key in ("/JavaScript", "/EmbeddedFiles")):
        raise UnsafeDocumentError("PDF scripts and embedded files are not allowed")
    if any(key in root for key in ("/OpenAction", "/AA")):
        raise UnsafeDocumentError("PDF automatic actions are not allowed")


def _decode_text(data: bytes) -> str:
    if b"\x00" in data:
        raise UnsafeDocumentError("Text document contains NUL bytes")
    for encoding in ("utf-8-sig", "utf-8", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnsafeDocumentError("Text document must use UTF-8 or UTF-16 encoding")


def validate_document(
    *,
    file_name: str,
    data: bytes,
    claimed_mime_type: str | None = None,
) -> ValidatedDocument:
    settings = get_settings()
    if not data:
        raise UnsafeDocumentError("Document is empty")
    if len(data) > settings.knowledge_max_upload_bytes:
        raise UnsafeDocumentError("Document exceeds the Knowledge upload limit")
    clean_name = sanitize_document_filename(file_name)
    extension = Path(clean_name).suffix.casefold()
    format_info = _FORMAT_BY_EXTENSION.get(extension)
    if format_info is None:
        raise UnsafeDocumentError(f"Unsupported Knowledge document extension: {extension}")
    format_name, canonical_mime = format_info
    claimed = (claimed_mime_type or "").split(";", 1)[0].strip().casefold()
    accepted_mimes = {canonical_mime.casefold()}
    if format_name in {"text", "markdown", "csv", "html", "json"}:
        accepted_mimes.add("text/plain")
    if claimed not in _GENERIC_MIMES and claimed not in accepted_mimes:
        raise UnsafeDocumentError("Claimed MIME type does not match the document extension")

    if format_name == "pdf":
        if not data.startswith(b"%PDF-"):
            raise UnsafeDocumentError("PDF magic bytes are missing")
        _validate_pdf(data)
    elif format_name in {"docx", "pptx", "xlsx"}:
        if not data.startswith(b"PK"):
            raise UnsafeDocumentError("Office document magic bytes are missing")
        _validate_ooxml_archive(data, extension)
    else:
        text = _decode_text(data)
        if format_name == "json":
            try:
                json.loads(text)
            except (TypeError, ValueError) as exc:
                raise UnsafeDocumentError("JSON document is invalid") from exc

    return ValidatedDocument(
        file_name=clean_name,
        extension=extension,
        mime_type=canonical_mime,
        size_bytes=len(data),
        format_name=format_name,
    )


def _normalize_extracted_text(value: str) -> str:
    value = unicodedata.normalize("NFC", (value or "").replace("\r\n", "\n").replace("\r", "\n"))
    value = value.replace("\x00", "")
    value = "".join(
        character for character in value if character in {"\n", "\t"} or unicodedata.category(character) != "Cc"
    )
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{4,}", "\n\n\n", value)
    return value.strip()


class _VisibleHTMLTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._hidden_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        del attrs
        if tag.casefold() in {"script", "style", "noscript", "svg"}:
            self._hidden_depth += 1
        elif not self._hidden_depth and tag.casefold() in {
            "p",
            "div",
            "br",
            "li",
            "tr",
            "h1",
            "h2",
            "h3",
            "h4",
            "h5",
            "h6",
        }:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.casefold() in {"script", "style", "noscript", "svg"} and self._hidden_depth:
            self._hidden_depth -= 1
        elif not self._hidden_depth:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._hidden_depth:
            self.parts.append(data)


def _bounded_segments(segments: list[ParsedSegment]) -> tuple[ParsedSegment, ...]:
    maximum = get_settings().knowledge_max_document_characters
    total = sum(len(segment.text) for segment in segments)
    if total > maximum:
        raise UnsafeDocumentError("Extracted document text exceeds the allowed limit")
    return tuple(segment for segment in segments if segment.text.strip())


def _ocr_dpi_candidates(preferred_dpi: int) -> tuple[int, ...]:
    base = min(400, max(100, preferred_dpi))
    ordered: list[int] = []
    for dpi in (base, min(base, 150), 120):
        if dpi not in ordered:
            ordered.append(dpi)
    return tuple(ordered)


def _ocr_pdf_page_text(
    *,
    data: bytes,
    page_number: int,
    settings,
    ocr_document_holder: list,
) -> str:
    """Render one PDF page and OCR it with memory-conscious fallbacks.

    The rasteriser is PDFium (``pypdfium2``, BSD-3 / Apache-2.0). It replaced
    PyMuPDF, which is AGPL or a commercial Artifex licence, and which this
    module used for exactly this one job: a grayscale bitmap of a page at a
    chosen DPI for ``pytesseract`` to read. Text extraction was always
    ``pypdf``. Measured before the swap on Latin, Persian and mixed scans:
    identical OCR text, identical peak memory, a few tens of milliseconds more
    per page against a multi-second tesseract pass.
    """
    import gc

    import pypdfium2 as pdfium
    import pytesseract

    last_error: Exception | None = None
    for attempt, dpi in enumerate(_ocr_dpi_candidates(settings.knowledge_ocr_dpi)):
        # Re-open periodically / after failures so renderer caches do not accumulate.
        if ocr_document_holder[0] is None or attempt > 0:
            if ocr_document_holder[0] is not None:
                ocr_document_holder[0].close()
                ocr_document_holder[0] = None
                gc.collect()
            ocr_document_holder[0] = pdfium.PdfDocument(data)
        page = None
        bitmap = None
        try:
            page = ocr_document_holder[0][page_number - 1]
            width, height = page.get_size()
            scale = dpi / 72
            if (width * scale) * (height * scale) > 30_000_000:
                raise UnsafeDocumentError("PDF page is too large for safe OCR")
            bitmap = page.render(scale=scale, grayscale=True)
            image = bitmap.to_pil().convert("L")
            try:
                ocr_text = pytesseract.image_to_string(
                    image,
                    lang=settings.knowledge_ocr_languages,
                    config="--oem 1 --psm 6",
                    timeout=settings.knowledge_ocr_page_timeout_seconds,
                )
            finally:
                image.close()
                del image
                gc.collect()
            return _normalize_extracted_text(ocr_text)
        except UnsafeDocumentError:
            raise
        except Exception as exc:  # noqa: BLE001 - retry lower DPI / remount doc
            last_error = exc
            if ocr_document_holder[0] is not None:
                ocr_document_holder[0].close()
                ocr_document_holder[0] = None
            gc.collect()
        finally:
            if bitmap is not None:
                bitmap.close()
            if page is not None:
                page.close()
    detail = ""
    if last_error is not None:
        detail = f": {type(last_error).__name__}: {last_error}"
    raise UnsafeDocumentError(f"OCR failed for PDF page {page_number}{detail}") from last_error


def parse_document(data: bytes, document: ValidatedDocument) -> ParsedDocument:  # noqa: C901 -- Phase 4 split; complexity must not grow
    segments: list[ParsedSegment] = []
    metadata: dict = {"format": document.format_name}
    if document.format_name == "pdf":
        reader = PdfReader(BytesIO(data), strict=True)
        metadata["page_count"] = len(reader.pages)
        settings = get_settings()
        ocr_document_holder: list = [None]
        ocr_pages: list[int] = []
        try:
            for page_number, page in enumerate(reader.pages, start=1):
                try:
                    text = _normalize_extracted_text(page.extract_text() or "")
                except Exception:  # noqa: BLE001 -- falls back to a safe default value
                    text = ""
                if len(text) < settings.knowledge_ocr_min_text_characters:
                    if page_number > settings.knowledge_ocr_max_pages:
                        raise UnsafeDocumentError("Scanned PDF exceeds the OCR page limit")
                    try:
                        normalized_ocr = _ocr_pdf_page_text(
                            data=data,
                            page_number=page_number,
                            settings=settings,
                            ocr_document_holder=ocr_document_holder,
                        )
                        if len(normalized_ocr) > len(text):
                            text = normalized_ocr
                            ocr_pages.append(page_number)
                    except UnsafeDocumentError:
                        raise
                    except Exception as exc:
                        if settings.knowledge_ocr_required:
                            raise UnsafeDocumentError(
                                f"OCR failed for PDF page {page_number}: {type(exc).__name__}: {exc}"
                            ) from exc
                if text:
                    segments.append(ParsedSegment(text=text, page_number=page_number))
        finally:
            if ocr_document_holder[0] is not None:
                ocr_document_holder[0].close()
        metadata["ocr_pages"] = ocr_pages
    elif document.format_name == "docx":
        word = WordDocument(BytesIO(data))
        current_heading: str | None = None
        buffer: list[str] = []
        for paragraph in word.paragraphs:
            text = _normalize_extracted_text(paragraph.text)
            if not text:
                continue
            style_name = (paragraph.style.name if paragraph.style else "").casefold()
            if style_name.startswith("heading"):
                if buffer:
                    segments.append(
                        ParsedSegment(
                            text="\n\n".join(buffer),
                            section=current_heading,
                        )
                    )
                    buffer = []
                current_heading = text
            else:
                buffer.append(text)
        for table in word.tables:
            for row in table.rows:
                cells = [_normalize_extracted_text(cell.text) for cell in row.cells]
                buffer.append(" | ".join(cell for cell in cells if cell))
        if buffer:
            segments.append(ParsedSegment(text="\n\n".join(buffer), section=current_heading))
    elif document.format_name == "pptx":
        presentation = Presentation(BytesIO(data))
        metadata["slide_count"] = len(presentation.slides)
        for slide_number, slide in enumerate(presentation.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = _normalize_extracted_text(shape.text)
                    if text:
                        texts.append(text)
            if texts:
                segments.append(
                    ParsedSegment(
                        text="\n\n".join(texts),
                        page_number=slide_number,
                        section=f"Slide {slide_number}",
                    )
                )
    elif document.format_name == "xlsx":
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        metadata["sheet_count"] = len(workbook.sheetnames)
        try:
            for sheet in workbook.worksheets:
                rows: list[str] = []
                for values in sheet.iter_rows(values_only=True):
                    cells = [_normalize_extracted_text(str(value)) for value in values if value is not None]
                    if cells:
                        rows.append(" | ".join(cells))
                if rows:
                    segments.append(
                        ParsedSegment(
                            text="\n".join(rows),
                            section=sheet.title,
                        )
                    )
        finally:
            workbook.close()
    else:
        text = _decode_text(data)
        if document.format_name == "html":
            parser = _VisibleHTMLTextParser()
            parser.feed(text)
            text = "".join(parser.parts)
        elif document.format_name == "json":
            text = json.dumps(
                json.loads(text),
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            )
        elif document.format_name == "csv":
            reader = csv.reader(StringIO(text))
            text = "\n".join(" | ".join(_normalize_extracted_text(cell) for cell in row) for row in reader)
        normalized = _normalize_extracted_text(text)
        if normalized:
            segments.append(ParsedSegment(text=normalized))

    bounded = _bounded_segments(segments)
    combined = "\n".join(segment.text for segment in bounded)
    security_flags = tuple(
        f"prompt_injection_pattern:{index + 1}"
        for index, pattern in enumerate(_PROMPT_INJECTION_PATTERNS)
        if pattern.search(combined)
    )
    return ParsedDocument(
        segments=bounded,
        character_count=sum(len(segment.text) for segment in bounded),
        metadata=metadata,
        security_flags=security_flags,
    )
