"""Unit tests for chat export to Word (.docx)."""

import io

from docx import Document

from app.services.chat_docx_service import (
    ChatExportError,
    render_chat_docx,
)
from app.services.chat_export_service import build_download_content_disposition


def _doc(content: str, title: str | None = None) -> Document:
    return Document(io.BytesIO(render_chat_docx(content=content, title=title)))


def test_docx_rejects_empty():
    try:
        render_chat_docx(content="   ")
        raise AssertionError("expected ChatExportError")
    except ChatExportError:
        pass


def test_docx_heading_and_paragraph():
    doc = _doc("# Title\n\nThis is a paragraph.")
    paras = doc.paragraphs
    # first paragraph is the heading
    assert any("Title" in p.text for p in paras)
    assert any("paragraph" in p.text for p in paras)


def test_docx_table_created():
    md = "| نام | سن |\n| --- | --- |\n| علی | ۳۰ |\n| سara | ۲۵ |"
    doc = _doc(md)
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert t.rows[0].cells[0].text.strip() == "نام"
    assert t.rows[1].cells[0].text.strip() == "علی"
    assert t.rows[2].cells[1].text.strip() == "۲۵"


def test_docx_code_block_preserved():
    md = "```python\nprint('hi')\n```"
    doc = _doc(md)
    code_text = "\n".join(p.text for p in doc.paragraphs)
    assert "print('hi')" in code_text


def test_docx_unordered_list():
    md = "- one\n- two\n- three"
    doc = _doc(md)
    texts = [p.text for p in doc.paragraphs]
    assert "one" in texts and "two" in texts and "three" in texts


def test_docx_ordered_list():
    md = "1. first\n2. second"
    doc = _doc(md)
    texts = [p.text for p in doc.paragraphs]
    assert "first" in texts and "second" in texts


def test_docx_persian_text_present():
    md = "این یک متن فارسی است.\n\n## بخش دوم\n\nمحتوای بیشتر."
    doc = _doc(md)
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "فارسی" in joined
    assert "بخش دوم" in joined


def test_docx_title_added_when_provided():
    doc = _doc("body text", title="گزارش فروش")
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "گزارش فروش" in joined


def test_docx_raw_html_treated_as_text():
    # Raw HTML must NOT be rendered/evaluated; it should appear as literal text.
    md = "<script>alert(1)</script>"
    doc = _doc(md)
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "alert(1)" in joined
    # No actual script execution possible (python-docx doesn't run code), but
    # ensure the text is preserved literally.


def test_docx_bold_inline():
    md = "This is **bold** text."
    doc = _doc(md)
    runs = [r for p in doc.paragraphs for r in p.runs]
    bold_runs = [r for r in runs if r.bold]
    assert any(r.text == "bold" for r in bold_runs)


def test_content_disposition_docx_extension():
    cd = build_download_content_disposition("My Report", "docx")
    assert 'filename="MyReport.docx"' in cd
    assert "filename*=UTF-8''" in cd
    assert cd.endswith(".docx")


def test_content_disposition_docx_persian_latin1_safe():
    cd = build_download_content_disposition("تحلیل فروش", "docx")
    cd.encode("latin-1")  # must not raise
    assert 'filename="chat-export.docx"' in cd
    assert "تحلیل" not in cd
