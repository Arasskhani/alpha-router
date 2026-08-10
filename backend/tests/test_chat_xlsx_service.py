"""Unit tests for chat export to Excel (.xlsx)."""

import io

from openpyxl import load_workbook

from app.services.chat_export_service import build_download_content_disposition
from app.services.chat_xlsx_service import ChatExportError, render_chat_xlsx


def _wb(content: str, title: str | None = None):
    return load_workbook(io.BytesIO(render_chat_xlsx(content=content, title=title)))


def test_xlsx_rejects_empty():
    try:
        render_chat_xlsx(content="   ")
        raise AssertionError("expected ChatExportError")
    except ChatExportError:
        pass


def test_xlsx_markdown_table():
    md = "| نام | سن |\n| --- | --- |\n| علی | ۳۰ |\n| سara | ۲۵ |"
    wb = _wb(md, title="Sales")
    ws = wb.active
    assert ws.title == "Sales"
    assert ws["A1"].value == "نام"
    assert ws["B1"].value == "سن"
    assert ws["A2"].value == "علی"
    assert ws["B3"].value == "۲۵"


def test_xlsx_csv_fence_fallback():
    md = "```csv\nname,value\na,1\nb,2\n```"
    wb = _wb(md)
    ws = wb.active
    assert ws["A1"].value == "name"
    assert ws["B2"].value == "1"


def test_xlsx_plain_text_fallback():
    wb = _wb("hello\nworld")
    ws = wb.active
    assert ws["A1"].value == "Content"
    assert ws["A2"].value == "hello"
    assert ws["A3"].value == "world"


def test_xlsx_content_disposition():
    cd = build_download_content_disposition("My Report", "xlsx")
    assert 'filename="MyReport.xlsx"' in cd
    assert ".xlsx" in cd
