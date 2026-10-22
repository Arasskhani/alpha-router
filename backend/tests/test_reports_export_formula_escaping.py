"""Report exports (all reports): text a spreadsheet would run as a formula is written as text."""

from __future__ import annotations

import csv
import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.services import reports_service

DANGEROUS = ['=HYPERLINK("http://x")', "+1+1", "-2+3", "@SUM(A1)", "\tlead", "\rlead"]


def _frame() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "name": [*DANGEROUS, "plain", None],
            "count": [-5, 3, 0, 7, 1, 2, -1, 4],
            "share": [-0.5, 1.5, 0.0, 2.0, 0.1, 0.2, 0.3, 0.4],
        }
    )


def test_csv_text_cells_are_escaped_and_numbers_kept():
    content, media, name = reports_service.export_dataframe(_frame(), "csv", "usage")
    assert media == "text/csv" and name == "usage.csv"
    rows = list(csv.reader(io.StringIO(content.decode("utf-8"), newline="")))
    assert rows[0] == ["name", "count", "share"]
    assert [r[0] for r in rows[1:7]] == ["'" + v for v in DANGEROUS]
    assert rows[7][0] == "plain" and rows[8][0] == ""
    assert [r[1] for r in rows[1:]] == ["-5", "3", "0", "7", "1", "2", "-1", "4"]
    assert rows[1][2] == "-0.5"


@pytest.mark.parametrize("fmt", ["xlsx", "xls"])
def test_xlsx_text_cells_are_strings_and_numbers_stay_numbers(fmt):
    content, _media, name = reports_service.export_dataframe(_frame(), fmt, "usage")
    assert name == "usage.xlsx"
    sheet = load_workbook(io.BytesIO(content)).active
    names = [sheet.cell(row=r, column=1) for r in range(2, 8)]
    assert [c.value for c in names] == ["'" + v for v in DANGEROUS]
    assert all(c.data_type == "s" for c in names)  # never "f" (a formula)
    assert sheet.cell(row=2, column=2).value == -5 and sheet.cell(row=2, column=2).data_type == "n"
    assert sheet.cell(row=2, column=3).value == -0.5


def test_a_header_that_looks_like_a_formula_is_escaped_too():
    content, _media, _name = reports_service.export_dataframe(pd.DataFrame({"=cmd": ["x"]}), "csv")
    assert content.decode("utf-8").splitlines()[0] == "'=cmd"


def test_the_caller_frame_is_left_as_it_was():
    frame = _frame()
    reports_service.export_dataframe(frame, "csv")
    assert frame["name"].tolist()[:6] == DANGEROUS
